"""Validate small project documents and extract text where needed."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
MAX_DOCUMENTS_PER_PROJECT = 3


def read_document(filename: str, content: bytes) -> tuple[str, str | None]:
    if not content:
        raise HTTPException(422, "Document is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Document exceeds the 8 MB limit")

    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise HTTPException(422, "Invalid PDF file")
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted or not 1 <= len(reader.pages) <= 100:
                raise ValueError("Unsupported PDF")
        except Exception as exc:
            raise HTTPException(422, "PDF must be readable, unencrypted, and at most 100 pages") from exc
        return "application/pdf", None

    if extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 500 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
                    raise ValueError("Archive too large")
            document = Document(BytesIO(content))
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(line.strip() for line in lines if line.strip())
        except Exception as exc:
            raise HTTPException(422, "Invalid or oversized DOCX file") from exc
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif extension in {".txt", ".md"}:
        try:
            text = content.decode("utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "Text files must use UTF-8") from exc
        media_type = "text/markdown" if extension == ".md" else "text/plain"
    else:
        raise HTTPException(415, "Supported files: PDF, DOCX, TXT, MD")

    if not text:
        raise HTTPException(422, "Document contains no readable text")
    if len(text) > MAX_TEXT_CHARS:
        raise HTTPException(413, "Document text exceeds the 100,000 character limit")
    return media_type, text
